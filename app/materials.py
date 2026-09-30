"""Materials intelligence — BOQ materials, current prices from the account's supplier price lists,
estimated material cost, and the same material across the account's active tenders.

Trust rules (never invent):
- BOQ lines come only from the tender's own tables, each with its file, page and row.
- A price is shown only when an uploaded supplier price list has the same material; the supplier,
  the list date and the file are always shown. Otherwise the line is PRICE_UNAVAILABLE.
- Quantity × unit price is an ESTIMATE of material cost, never the tender price.
- Two lines are merged only when their normalised keys are identical (same family, sizes, ratings
  and unit). Same family and main size with other details missing or different = REQUIRES_REVIEW.
- A bulk scenario is shown only when a price list itself has a lower price for a larger quantity.
"""
from __future__ import annotations

import re
from collections import defaultdict
from datetime import datetime
from typing import Any, Dict, Iterable, List, Optional, Tuple

from app.pipeline.structured_data import map_boq_columns

# ---- normalisation
_AR_DIGITS = str.maketrans("٠١٢٣٤٥٦٧٨٩٫٬", "0123456789.,")

FAMILIES: List[Tuple[str, str]] = [  # order matters: specific before generic
    ("cable_termination", r"terminat|cable\s+joint|straight\s+joint|نهاي(?:ة|ات)\s+كابل|وصل(?:ة|ات)\s+كابل"),
    ("cable_tray", r"cable\s+(?:tray|ladder|trunking)|حامل\s+كابلات|مجرى\s+كابلات"),
    ("cable", r"\bcables?\b|كابل|كبل"),
    ("wire", r"\bwires?\b|conductor|سلك|أسلاك"),
    ("circuit_breaker", r"circuit\s+breaker|\bmccb\b|\bmcb\b|\bacb\b|\bvcb\b|\brccb\b|breaker|قاطع"),
    ("switchgear", r"switchgear|ring\s+main\s+unit|\brmu\b|\bgis\b|خلية|خلايا|وحدة\s+حلقية"),
    ("transformer", r"transformer|محول"),
    ("panel", r"\bpanel|switchboard|distribution\s+board|\bmdb\b|\bsmdb\b|\bdb\b|لوحة|لوحات"),
    ("conduit_pipe", r"conduit|\bpipes?\b|upvc|\bhdpe\b|ماسورة|مواسير|أنابيب"),
    ("light_fitting", r"luminaire|light\s+fitting|lighting\s+fixture|\blamp|flood\s*light|كشاف|وحدة\s+إضاءة|وحدات\s+إنارة"),
    ("earthing", r"earth(?:ing)?\s+rod|grounding|earthing|تأريض|أرضي"),
    ("relay", r"\brelays?\b|ريلاي|مرحل"),
    ("meter", r"\bmeters?\b(?!\s*(?:long|length))|عداد"),
    ("busduct", r"bus\s*duct|busway|bus\s*bar\s*trunking"),
]
_FAM = [(k, re.compile(p, re.IGNORECASE)) for k, p in FAMILIES]
_SIZE = re.compile(r"(\d{1,2})\s*(?:c(?:ore)?\s*)?[x×\*]\s*(\d{1,4}(?:\.\d)?)\s*(?:mm2|mm²|sq\.?\s?mm|mm|مم2|مم)?", re.IGNORECASE)
_SECTION = re.compile(r"(\d{1,4}(?:\.\d)?)\s*(?:mm2|mm²|sq\.?\s?mm|مم2)", re.IGNORECASE)
# "0.6/1 kV" and "1 kV" are the same cable rating; "11/0.4 kV" is an 11 kV transformer -> the higher value
_KV = re.compile(r"(\d{1,3}(?:\.\d{1,2})?)(?:\s*/\s*(\d{1,3}(?:\.\d{1,2})?))?\s*(?:kv\b|ك\.?\s?ف)", re.IGNORECASE)
# works lines (digging a cable trench) are not materials
_WORKS = re.compile(r"^\W*(?:excavat|backfill|trench|civil|dismantl|removal|testing|commissioning|حفر|ردم|فك|اختبار)",
                    re.IGNORECASE)
_AMP = re.compile(r"(\d{1,5})\s*(?:a|amp|amps|أمبير)\b", re.IGNORECASE)
_POLE = re.compile(r"\b([1-4])\s*(?:p|pole|poles)\b|\b(tp|dp|sp)\b", re.IGNORECASE)
_KVA = re.compile(r"(\d{1,6}(?:\.\d)?)\s*(kva|mva)\b", re.IGNORECASE)
_DIA = re.compile(r"(\d{1,4})\s*mm\b(?!\s*2)|(\d(?:\.\d+)?)\s*(?:\"|inch|in\b|بوصة)", re.IGNORECASE)
_KA = re.compile(r"(\d{1,3})\s*ka\b", re.IGNORECASE)
_WATT = re.compile(r"(\d{1,4})\s*(?:w|watt|وات)\b", re.IGNORECASE)
_TAGS = [("xlpe", r"xlpe"), ("pvc", r"\bpvc\b"), ("armoured", r"armou?r|\bswa\b|\bsta\b|مسلح"),
         ("cu", r"\bcu\b|copper|نحاس"), ("al", r"\bal\b|alumin|ألومنيوم|المونيوم"),
         ("led", r"\bled\b"), ("fire_resistant", r"fire\s*(?:resistant|rated)|\bfr\b|مقاوم\s+للحريق")]
_TAGS_RX = [(k, re.compile(p, re.IGNORECASE)) for k, p in _TAGS]

UNITS = {"m": ["m", "mtr", "meter", "metre", "meters", "metres", "lm", "l.m", "rm", "م", "م.ط", "متر", "مترطولي", "متر طولي"],
         "km": ["km", "كم"],
         "pcs": ["no", "nos", "no.", "nr", "pcs", "pc", "each", "ea", "unit", "عدد", "قطعة"],
         "set": ["set", "sets", "طقم", "مجموعة"], "lot": ["lot", "ls", "l.s", "lump sum", "مقطوعية"],
         "kg": ["kg", "كجم", "كيلو"], "ton": ["ton", "tons", "t", "طن"], "m2": ["m2", "sqm", "م2"],
         "m3": ["m3", "cum", "م3"]}
_UNIT = {a: k for k, al in UNITS.items() for a in al}


def norm_unit(u: Optional[str]) -> str:
    u = " ".join(str(u or "").translate(_AR_DIGITS).lower().replace("'", "").split())
    return _UNIT.get(u, _UNIT.get(u.rstrip("."), u))


def to_number(v: Any) -> Optional[float]:
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).translate(_AR_DIGITS).strip().replace(",", "").replace(" ", "")
    m = re.fullmatch(r"-?\d+(?:\.\d+)?", s)
    return float(m.group(0)) if m else None


def material_attrs(text: str) -> Dict[str, Any]:
    """Family + technical attributes read from a description (English / Arabic)."""
    t = " ".join(str(text or "").translate(_AR_DIGITS).split())
    fam = None if _WORKS.search(t) else next((k for k, rx in _FAM if rx.search(t)), None)
    a: Dict[str, Any] = {"family": fam}
    m = _SIZE.search(t)
    if m:
        a["size"] = f"{int(m.group(1))}x{float(m.group(2)):g}"
    else:
        s = _SECTION.search(t)
        if s:
            a["size"] = f"1x{float(s.group(1)):g}"
    kv = _KV.search(t)
    if kv:
        a["kv"] = f"{max(float(x) for x in kv.groups() if x):g}"
    amp = _AMP.search(t)
    if amp and fam not in ("cable", "wire"):
        a["amp"] = int(amp.group(1))
    p = _POLE.search(t)
    if p:
        a["poles"] = int(p.group(1)) if p.group(1) else {"sp": 1, "dp": 2, "tp": 3}[p.group(2).lower()]
    k = _KVA.search(t)
    if k:
        a["rating"] = f"{float(k.group(1)):g}{k.group(2).lower()}"
    w = _WATT.search(t)
    if w and fam == "light_fitting":
        a["watt"] = int(w.group(1))
    ka = _KA.search(t)
    if ka:
        a["ka"] = int(ka.group(1))
    if fam in ("conduit_pipe", "cable_tray", "earthing"):
        d = _DIA.search(t)
        if d:
            a["dia"] = f"{d.group(1)}mm" if d.group(1) else f"{d.group(2)}in"
    tags = {k for k, rx in _TAGS_RX if rx.search(t)}
    if "xlpe" in tags:
        tags.discard("pvc")  # PVC is then only the outer sheath
    tags = sorted(tags)
    if tags:
        a["tags"] = tags
    return a


MAIN_ATTR = {"cable": "size", "wire": "size", "cable_termination": "size", "circuit_breaker": "amp",
             "transformer": "rating", "conduit_pipe": "dia", "cable_tray": "dia", "switchgear": "kv", "light_fitting": "watt"}


def material_key(attrs: Dict[str, Any], unit: str) -> Optional[str]:
    """Identical keys = the same material. None when the family or its main size is unknown."""
    fam = attrs.get("family")
    if not fam:
        return None
    main = MAIN_ATTR.get(fam)
    if main and not attrs.get(main):
        return None
    parts = [fam] + [f"{k}={attrs[k]}" for k in ("size", "kv", "amp", "poles", "rating", "ka", "dia", "watt") if attrs.get(k)]
    if attrs.get("tags"):
        parts.append("tags=" + "+".join(attrs["tags"]))
    return "|".join(parts) + f"|unit={unit}"


def family_key(attrs: Dict[str, Any], unit: str) -> Optional[str]:
    """Family + main size + unit: lines sharing this but not the full key need a person to review."""
    fam = attrs.get("family")
    main = MAIN_ATTR.get(fam or "")
    if not fam or not main or not attrs.get(main):
        return None
    return f"{fam}|{main}={attrs[main]}|unit={unit}"


def label(attrs: Dict[str, Any]) -> str:
    fam = (attrs.get("family") or "item").replace("_", " ")
    bits = [attrs.get("size") and f"{attrs['size']} mm²", attrs.get("kv") and f"{attrs['kv']} kV",
            attrs.get("amp") and f"{attrs['amp']} A", attrs.get("poles") and f"{attrs['poles']}P",
            attrs.get("rating"), attrs.get("ka") and f"{attrs['ka']} kA", attrs.get("dia"), attrs.get("watt") and f"{attrs['watt']} W"]
    bits += [x.upper() if len(x) <= 5 else x.replace("_", " ") for x in attrs.get("tags") or []]
    return " ".join([fam.capitalize()] + [str(b) for b in bits if b])


# ---- BOQ lines from the tender's own tables
def _rows(text: str) -> List[List[str]]:
    return [[c.strip() for c in ln.split("|")] for ln in (text or "").splitlines() if ln.count("|") >= 2]


def boq_lines(sources: Iterable[Any]) -> List[Dict[str, Any]]:
    """Every table row with a description and a positive quantity, under a BOQ-like header row."""
    out = []
    for s in sources:
        rows = _rows(s.text)
        header = None
        for i, cells in enumerate(rows):
            m = map_boq_columns(cells)
            slots = set(m.values())
            if {"description", "quantity"} <= slots:
                header = {}
                for j, col in enumerate(cells):  # first column per slot ("Qty" before "Amount")
                    if col in m:
                        header.setdefault(m[col], j)
                continue
            if not header:
                continue
            get = lambda slot: cells[header[slot]] if slot in header and header[slot] < len(cells) else ""
            qty = to_number(get("quantity"))
            desc = " ".join(get("description").split())
            if not qty or qty <= 0 or len(desc) < 4:
                continue
            unit = norm_unit(get("unit"))
            if unit == "km":
                qty, unit = qty * 1000, "m"
            attrs = material_attrs(desc)
            out.append({"document": s.source_document, "page": s.page_number, "row": i + 1,
                        "item": get("item") or None, "description": desc[:300], "unit": unit, "quantity": qty,
                        "attrs": attrs, "name": label(attrs) if attrs.get("family") else desc[:80],
                        "key": material_key(attrs, unit), "family_key": family_key(attrs, unit)})
    return out


_CACHE: Dict[str, Any] = {}


def tender_boq(tender_id: str) -> List[Dict[str, Any]]:
    from app.pipeline.checkpoint import cache_dir
    from app.sections import sources_from_cache
    d = cache_dir(tender_id)
    stamp = tuple(sorted((f.name, f.stat().st_mtime) for f in d.glob("extract-v*.json"))) if d.is_dir() else ()
    hit = _CACHE.get(tender_id)
    if hit and hit[0] == stamp:
        return hit[1]
    res = boq_lines(sources_from_cache(tender_id))
    if len(_CACHE) > 32:
        _CACHE.clear()
    _CACHE[tender_id] = (stamp, res)
    return res


# ---- supplier price lists
PRICE_ALIASES = {
    "description": {"description", "material", "item description", "product", "الوصف", "البيان", "الصنف", "المادة"},
    "unit": {"unit", "uom", "الوحدة"},
    "price": {"price", "unit price", "rate", "net price", "سعر", "السعر", "سعر الوحدة"},
    "currency": {"currency", "cur", "العملة"},
    "min_qty": {"min qty", "minimum quantity", "min quantity", "from qty", "moq", "الحد الأدنى", "من كمية"},
}


def _price_columns(cells: List[str]) -> Dict[str, int]:
    out = {}
    for i, c in enumerate(cells):
        n = " ".join(str(c or "").lower().split())
        for slot, al in PRICE_ALIASES.items():
            if slot not in out and (n in al or any(" " in a and a in n for a in al)):
                out[slot] = i
                break
    return out


def parse_price_rows(rows: List[List[Any]], default_currency: str) -> List[Dict[str, Any]]:
    """Rows of a supplier sheet -> price items. Rows without a number price are skipped, never guessed."""
    cols, out = None, []
    for r in rows:
        cells = ["" if c is None else str(c) for c in r]
        if cols is None:
            c = _price_columns(cells)
            if "description" in c and "price" in c:
                cols = c
            continue
        get = lambda k: cells[cols[k]] if k in cols and cols[k] < len(cells) else ""
        price = to_number(get("price"))
        desc = " ".join(get("description").split())
        if not price or price <= 0 or not desc:
            continue
        unit = norm_unit(get("unit"))
        if unit == "km":
            price, unit = price / 1000, "m"
        attrs = material_attrs(desc)
        out.append({"description": desc[:300], "unit": unit, "price": price,
                    "currency": (get("currency") or default_currency).strip().upper()[:3],
                    "min_qty": to_number(get("min_qty")) or 0, "key": material_key(attrs, unit)})
    return out


def read_sheet(path: str) -> List[List[Any]]:
    if path.lower().endswith(".csv"):
        import csv
        with open(path, encoding="utf-8-sig", newline="") as f:
            return [row for row in csv.reader(f)]
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        return [list(r) for ws in wb.worksheets for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


STALE_DAYS = 90


def price_index(price_items: List[Dict[str, Any]]) -> Dict[str, List[Dict[str, Any]]]:
    """key -> price items, newest list first, then the base (min_qty 0) price first."""
    idx: Dict[str, List[Dict[str, Any]]] = defaultdict(list)
    for p in price_items:
        if p.get("key"):
            idx[p["key"]].append(p)
    for v in idx.values():
        v.sort(key=lambda p: (-(p["price_date"].timestamp() if p.get("price_date") else 0), p.get("min_qty") or 0))
    return idx


def _price_view(p: Dict[str, Any], now: datetime) -> Dict[str, Any]:
    age = (now - p["price_date"]).days if p.get("price_date") else None
    return {"price": p["price"], "currency": p["currency"], "unit": p["unit"], "supplier": p["supplier"],
            "price_date": p["price_date"].date().isoformat() if p.get("price_date") else None,
            "source_file": p.get("filename"), "list_id": p.get("list_id"), "age_days": age,
            "stale": age is not None and age > STALE_DAYS}


def current_price(key: Optional[str], idx: Dict[str, List[Dict[str, Any]]], qty: float = 0,
                  now: Optional[datetime] = None) -> Optional[Dict[str, Any]]:
    """The newest list's base price for this exact material. None = PRICE_UNAVAILABLE."""
    if not key or key not in idx:
        return None
    now = now or datetime.utcnow()
    items = idx[key]
    newest = items[0].get("list_id")
    base = next((p for p in items if p.get("list_id") == newest and not p.get("min_qty")), items[0])
    return _price_view(base, now)


def bulk_price(key: Optional[str], idx: Dict[str, List[Dict[str, Any]]], qty: float,
               now: Optional[datetime] = None) -> Optional[Dict[str, Any]]:
    """A lower price the same supplier list quotes for this larger quantity (a real price break)."""
    if not key or key not in idx:
        return None
    now = now or datetime.utcnow()
    items = idx[key]
    newest = items[0].get("list_id")
    tiers = [p for p in items if p.get("list_id") == newest and (p.get("min_qty") or 0) > 0 and p["min_qty"] <= qty]
    if not tiers:
        return None
    best = min(tiers, key=lambda p: p["price"])
    return dict(_price_view(best, now), min_qty=best["min_qty"])


def cost_view(lines: List[Dict[str, Any]], idx, currency: Optional[str]) -> Dict[str, Any]:
    """Priced lines, the unavailable ones, and totals per currency (never converted)."""
    now = datetime.utcnow()
    out, totals, unavailable = [], defaultdict(float), 0
    for ln in lines:
        p = current_price(ln["key"], idx, ln["quantity"], now)
        item = dict(ln, price=p, status="PRICED" if p else "PRICE_UNAVAILABLE")
        if p:
            item["cost"] = round(ln["quantity"] * p["price"], 2)
            totals[p["currency"]] += item["cost"]
        else:
            unavailable += 1
        out.append(item)
    return {"lines": out, "totals": [{"currency": c, "amount": round(v, 2)} for c, v in totals.items()],
            "unavailable": unavailable, "priced": len(out) - unavailable, "tender_currency": currency,
            "label": "ESTIMATE"}


def aggregate(tender_lines: Dict[str, List[Dict[str, Any]]], titles: Dict[str, str], idx) -> Dict[str, Any]:
    """The same material in two or more active tenders (identical keys), and possible matches to review."""
    now = datetime.utcnow()
    by_key: Dict[str, Dict[str, Any]] = {}
    fam: Dict[str, Dict[str, Any]] = defaultdict(lambda: {"keys": set(), "tenders": set(), "lines": []})
    for tid, lines in tender_lines.items():
        for ln in lines:
            if ln["key"]:
                g = by_key.setdefault(ln["key"], {"key": ln["key"], "name": ln["name"], "unit": ln["unit"],
                                                  "tenders": defaultdict(float), "sources": []})
                g["tenders"][tid] += ln["quantity"]
                g["sources"].append({"tender_id": tid, "document": ln["document"], "page": ln["page"],
                                     "row": ln["row"], "description": ln["description"], "quantity": ln["quantity"]})
            if ln["family_key"]:
                f = fam[ln["family_key"]]
                f["keys"].add(ln["key"] or ln["description"])
                f["tenders"].add(tid)
                f["lines"].append({"tender_id": tid, "description": ln["description"], "quantity": ln["quantity"],
                                   "unit": ln["unit"], "document": ln["document"], "page": ln["page"]})
    bulk = []
    for g in by_key.values():
        if len(g["tenders"]) < 2:
            continue
        total = sum(g["tenders"].values())
        p = current_price(g["key"], idx, total, now)
        b = bulk_price(g["key"], idx, total, now)
        row = {"key": g["key"], "name": g["name"], "unit": g["unit"], "total_quantity": total,
               "tenders": [{"tender_id": t, "title": titles.get(t), "quantity": q} for t, q in g["tenders"].items()],
               "sources": g["sources"][:20], "price": p, "status": "PRICED" if p else "PRICE_UNAVAILABLE"}
        if p:
            row["estimated_cost"] = round(total * p["price"], 2)
        if p and b and b["currency"] == p["currency"] and b["price"] < p["price"]:
            row["bulk_scenario"] = {"price": b, "estimated_cost": round(total * b["price"], 2),
                                    "difference": round(total * (p["price"] - b["price"]), 2),
                                    "label": "ESTIMATED_SCENARIO"}
        bulk.append(row)
    bulk.sort(key=lambda r: (-len(r["tenders"]), -r["total_quantity"]))
    review = [{"family_key": k, "tenders": sorted(v["tenders"]), "lines": v["lines"][:12], "status": "REQUIRES_REVIEW"}
              for k, v in fam.items() if len(v["keys"]) > 1 and len(v["tenders"]) > 1]
    return {"bulk": bulk, "review": review}
