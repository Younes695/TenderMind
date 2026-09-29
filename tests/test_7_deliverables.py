"""Stage 7 — submission checklist, audit trail, compliance matrix, decision pack."""
import io
import uuid

import pytest
from fastapi.testclient import TestClient

from app.checklist import build, is_bid_item


@pytest.fixture(autouse=True, scope="module")
def _db_ready():
    from app.database import init_db
    init_db()


@pytest.mark.parametrize("text,expected", [
    ("Bidder shall provide duly signed and stamped softcopy of the complete tender documents.", True),
    ("The Bidder's Commercial Proposal shall also be accompanied with an exact copy of Annexure VII.", True),
    ("Completely filled-in Data Schedules, duly signed, stamped & dated.", True),
    ("The CONTRACTOR shall submit O&M manuals before energization.", False),           # execution submittal
    ("Contractor shall submit the training schedule for COMPANY approval.", False),
    ("B = Original Total Material Bid Price for Power Transformer", False),            # formula
    ("1.7 “CONTRACTOR” means the firm or consortium submitting the bid.", False),  # definition
    ("Only firms prequalified and invited may submit bids.", False),                   # eligibility, not an item
])
def test_bid_items(text, expected):
    assert is_bid_item(text) is expected


def test_build_merges_near_duplicates():
    reqs = [{"summary": "a", "source_text": "Completely filled-in Data Schedules given in the SOW/TS, duly signed, stamped.",
             "source_document": "ITB.doc", "page_number": 1},
            {"summary": "b", "source_text": "Completely filled-in Data Schedules given in the SOW/TS duly signed and stamped",
             "source_document": "SOW.pdf", "page_number": 103},
            {"summary": "c", "source_text": "Bidder shall provide QA/QC Plan for the complete project.",
             "source_document": "ITB.doc", "page_number": 1}]
    items = build(reqs)
    assert len(items) == 2 and items[0]["file"] == "ITB.doc"


def _tender_with_analysis(c):
    from app.database import SessionLocal
    from app.models import TenderAnalysis
    tid = f"S7-{uuid.uuid4().hex[:6]}"
    c.post("/api/tenders", json={"id": tid, "title": "Riyadh 132kV Substation"})
    db = SessionLocal()
    db.add(TenderAnalysis(id=f"AN-{uuid.uuid4().hex[:8]}", tender_id=tid, analysis_version="t", status="COMPLETED",
                          tender={}, evidence=[], documents=[], deadlines=[], commercial=None, risks=[],
                          derived_features={},
                          requirements=[{"requirement_id": "R1", "summary": "QA/QC plan", "category": "SUBMISSION",
                                         "mandatory": True, "source_document": "ITB.doc", "page_number": 2,
                                         "source_text": "Bidder shall provide QA/QC Plan with the proposal."}]))
    db.commit(); db.close()
    return tid


def test_checklist_state_and_audit_trail():
    from app.main import app
    c = TestClient(app)
    tid = _tender_with_analysis(c)
    cl = c.get(f"/api/tenders/{tid}/checklist").json()
    assert cl["progress"] == {"TODO": 1, "READY": 0, "NOT_APPLICABLE": 0, "total": 1}
    key = cl["items"][0]["key"]
    r = c.put(f"/api/tenders/{tid}/checklist/{key}", json={"status": "ready", "assignee": "Mona"}).json()
    assert r["status"] == "READY" and r["assignee"] == "Mona" and r["page"] == 2
    assert c.put(f"/api/tenders/{tid}/checklist/{key}", json={"status": "later"}).status_code == 422
    assert c.put(f"/api/tenders/{tid}/checklist/nope", json={"status": "READY"}).status_code == 404
    c.put(f"/api/tenders/{tid}/votes", json={"member_name": "Sara", "department": "Finance", "vote": "REJECT"})
    c.put(f"/api/tenders/{tid}/outcome", json={"outcome": "SUBMITTED"})
    pack = c.get(f"/api/tenders/{tid}/decision-pack").json()
    actions = [e["action"] for e in pack["audit"]]
    assert {"checklist_item", "vote_saved", "outcome_set"} <= set(actions)
    assert pack["checklist"]["progress"]["READY"] == 1
    assert list(pack)[2] == "certifications"  # certifications come first after the header


def test_compliance_matrix_xlsx_both_languages():
    from openpyxl import load_workbook
    from app.main import app
    c = TestClient(app)
    tid = _tender_with_analysis(c)
    for lang, sheet in (("en", "Matrix"), ("ar", "المصفوفة")):
        r = c.get(f"/api/tenders/{tid}/compliance-matrix.xlsx", params={"lang": lang})
        assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
        wb = load_workbook(io.BytesIO(r.content))
        assert sheet in wb.sheetnames
        if lang == "ar":
            assert wb[sheet].sheet_view.rightToLeft
        summary = wb[wb.sheetnames[0]]
        assert any("Method" in str(x.value) or "طريقة" in str(x.value) for x in summary["A"])
    assert c.get("/api/tenders/NOPE-1/compliance-matrix.xlsx").status_code == 404
